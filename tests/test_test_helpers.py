"""確認測試輔助工具本身可靠，後續模組的測試才有意義。"""

import docx
import numpy as np
import openpyxl
import pymupdf
import pytest

from rag_chatbot.embedder import EmbedderError
from tests.builders import make_csv, make_docx, make_pdf, make_xlsx
from tests.fakes import FakeEmbedder


def _cosine(a, b):
    return float(a @ b / (np.linalg.norm(a) * np.linalg.norm(b)))


def test_fake_embedder_similarity():
    embedder = FakeEmbedder()
    v = embedder.embed(["洗衣機保固兩年", "保固多久？", "明天天氣如何"])
    assert v.shape == (3, embedder.dim)
    assert v.dtype == np.float32
    assert _cosine(v[0], v[1]) > 0.9
    assert _cosine(v[0], v[2]) < 0.1


def test_fake_embedder_no_zero_vector():
    v = FakeEmbedder().embed(["沒有任何關鍵字"])
    assert np.linalg.norm(v[0]) > 0


def test_fake_embedder_records_calls_and_fails():
    embedder = FakeEmbedder(fail_on_call=2)
    embedder.embed(["a"])
    with pytest.raises(EmbedderError, match="模擬失敗"):
        embedder.embed(["b", "c"])
    assert embedder.calls == [["a"], ["b", "c"]]


def test_make_pdf(tmp_path):
    path = make_pdf(tmp_path / "t.pdf", ["第一頁：產品簡介", "", "第三頁：保固 2 年"])
    with pymupdf.open(path) as doc:
        texts = [page.get_text() for page in doc]
    assert len(texts) == 3
    assert "產品簡介" in texts[0]
    assert texts[1].strip() == ""
    assert "保固 2 年" in texts[2]


def test_make_docx(tmp_path):
    path = make_docx(
        tmp_path / "t.docx",
        [("heading", 1, "工作經歷"), ("para", "後端工程師"), ("table", [["年份", "職稱"], ["2024", "工程師"]])],
    )
    document = docx.Document(path)
    paragraphs = [(p.style.name, p.text) for p in document.paragraphs]
    assert ("Heading 1", "工作經歷") in paragraphs
    assert ("Normal", "後端工程師") in paragraphs
    assert document.tables[0].cell(1, 1).text == "工程師"


def test_make_docx_unknown_item(tmp_path):
    with pytest.raises(ValueError):
        make_docx(tmp_path / "t.docx", [("image", "x")])


@pytest.mark.parametrize("encoding", ["utf-8-sig", "cp950"])
def test_make_csv(tmp_path, encoding):
    path = make_csv(tmp_path / "t.csv", [["question", "answer"], ["健保卡遺失？", "補辦"]], encoding)
    assert path.read_bytes().decode(encoding).splitlines()[1] == "健保卡遺失？,補辦"


def test_make_xlsx(tmp_path):
    path = make_xlsx(tmp_path / "t.xlsx", [["question", "answer"], ["共用帳本？", "最多 5 人"]])
    rows = list(openpyxl.load_workbook(path).active.iter_rows(values_only=True))
    assert rows == [("question", "answer"), ("共用帳本？", "最多 5 人")]
