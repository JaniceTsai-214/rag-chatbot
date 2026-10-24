"""在測試執行時產生小型測試文件，不把二進位檔放進 Git（docs/04-testing.md §1）。"""

from __future__ import annotations

import csv
from collections.abc import Sequence
from pathlib import Path

import docx
import openpyxl
import pymupdf

# PyMuPDF 內建的繁體中文字型
PDF_FONT = "china-t"


def make_pdf(path: Path, pages: Sequence[str]) -> Path:
    """每個字串一頁；空字串產生空白頁。"""
    doc = pymupdf.open()
    for text in pages:
        page = doc.new_page()
        if text:
            page.insert_textbox(page.rect + (50, 50, -50, -50), text, fontname=PDF_FONT, fontsize=11)
    doc.save(path)
    doc.close()
    return path


def make_docx(path: Path, items: Sequence[tuple]) -> Path:
    """依序加入內容。

    ("heading", level, text)：標題（套用「Heading N」樣式）
    ("para", text)：一般段落
    ("table", rows)：表格，rows 為字串的二維清單
    """
    document = docx.Document()
    for item in items:
        kind = item[0]
        if kind == "heading":
            document.add_heading(item[2], level=item[1])
        elif kind == "para":
            document.add_paragraph(item[1])
        elif kind == "table":
            rows = item[1]
            table = document.add_table(rows=len(rows), cols=len(rows[0]))
            for r, row in enumerate(rows):
                for c, value in enumerate(row):
                    table.cell(r, c).text = value
        else:
            raise ValueError(f"未知的項目類型：{kind}")
    document.save(path)
    return path


def make_csv(path: Path, rows: Sequence[Sequence[str]], encoding: str = "utf-8-sig") -> Path:
    """第一列為欄位名稱。預設 UTF-8 含 BOM（Excel 另存 UTF-8 CSV 的格式）。"""
    with path.open("w", encoding=encoding, newline="") as f:
        csv.writer(f).writerows(rows)
    return path


def make_xlsx(path: Path, rows: Sequence[Sequence[str]]) -> Path:
    """第一列為欄位名稱。"""
    workbook = openpyxl.Workbook()
    sheet = workbook.active
    for row in rows:
        sheet.append(list(row))
    workbook.save(path)
    return path
