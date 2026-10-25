from pathlib import Path

import pymupdf
import pytest

from rag_chatbot.loaders import LoaderError, file_type, load
from rag_chatbot.models import Section
from tests.builders import make_csv, make_docx, make_pdf, make_xlsx


# ---------- 分派 ----------

@pytest.mark.parametrize(
    "name, expected",
    [("a.pdf", "pdf"), ("a.DOCX", "docx"), ("a.md", "markdown"),
     ("a.csv", "faq"), ("a.xlsx", "faq"), ("a.txt", None), ("a", None)],
)
def test_file_type(name, expected):
    assert file_type(Path(name)) == expected


def test_unsupported_format(tmp_path):
    path = tmp_path / "notes.txt"
    path.write_text("hello", encoding="utf-8")
    with pytest.raises(LoaderError, match="不支援的格式"):
        load(path)


def test_missing_file(tmp_path):
    with pytest.raises(LoaderError, match="找不到檔案"):
        load(tmp_path / "nope.md")


@pytest.mark.parametrize("name", ["empty.md", "empty.csv"])
def test_empty_text_file(tmp_path, name):
    (tmp_path / name).write_bytes(b"")
    with pytest.raises(LoaderError, match="沒有可擷取的文字"):
        load(tmp_path / name)


# ---------- PDF ----------

def test_pdf_pages(tmp_path):
    path = make_pdf(tmp_path / "手冊.pdf", ["產品簡介", "", "方案價格：進階版每月 60 元"])
    sections = load(path)
    assert [s.page for s in sections] == [1, 3]  # 空白頁略過，頁碼維持原本編號
    assert sections[0].text == "產品簡介"
    assert "60 元" in sections[1].text
    assert all(s.heading is None for s in sections)


def test_pdf_without_text_layer(tmp_path):
    path = make_pdf(tmp_path / "scan.pdf", ["", ""])
    with pytest.raises(LoaderError, match="沒有文字層"):
        load(path)


def test_pdf_broken(tmp_path):
    path = tmp_path / "bad.pdf"
    path.write_bytes(b"not a pdf at all")
    with pytest.raises(LoaderError, match="無法開啟 PDF"):
        load(path)


def test_pdf_password(tmp_path):
    doc = pymupdf.open()
    doc.new_page().insert_text((50, 72), "secret", fontname="helv")
    path = tmp_path / "locked.pdf"
    doc.save(path, encryption=pymupdf.PDF_ENCRYPT_AES_256, user_pw="u", owner_pw="o")
    with pytest.raises(LoaderError, match="密碼保護"):
        load(path)


# ---------- DOCX ----------

def test_docx_sections_by_heading(tmp_path):
    path = make_docx(tmp_path / "履歷.docx", [
        ("para", "王小明，後端工程師。"),
        ("heading", 1, "工作經歷"),
        ("heading", 2, "小算盤團隊"),
        ("para", "負責雲端同步功能。"),
        ("heading", 1, "專長"),
        ("para", "Python"),
        ("para", "SQL"),
        ("heading", 1, "空章節"),
    ])
    assert load(path) == [
        Section(text="王小明，後端工程師。", heading=None),
        Section(text="負責雲端同步功能。", heading="工作經歷 > 小算盤團隊"),
        Section(text="Python\nSQL", heading="專長"),
    ]


def test_docx_without_headings(tmp_path):
    path = make_docx(tmp_path / "a.docx", [("para", "第一段"), ("para", ""), ("para", "第二段")])
    assert load(path) == [Section(text="第一段\n\n第二段", heading=None)]


def test_docx_title_style(tmp_path):
    path = make_docx(tmp_path / "a.docx", [("heading", 0, "王小明"), ("heading", 1, "專長"), ("para", "Python")])
    assert load(path) == [Section(text="Python", heading="王小明 > 專長")]


def test_docx_table(tmp_path):
    path = make_docx(tmp_path / "a.docx", [
        ("heading", 1, "經歷"),
        ("para", "摘要"),
        ("table", [["年份", "職稱"], ["2024", "工程師"], ["", ""]]),
        ("para", "結尾"),
    ])
    assert load(path) == [Section(text="摘要\n年份 | 職稱\n2024 | 工程師\n結尾", heading="經歷")]


def test_docx_merged_cells(tmp_path):
    path = make_docx(tmp_path / "a.docx", [("table", [["方案", "價格"], ["進階版", "60 元"]])])
    import docx as python_docx

    document = python_docx.Document(path)
    table = document.tables[0]
    table.cell(0, 0).merge(table.cell(0, 1))
    document.save(path)
    assert load(path)[0].text.splitlines()[1] == "進階版 | 60 元"
    assert "|" not in load(path)[0].text.splitlines()[0]


def test_docx_broken(tmp_path):
    path = tmp_path / "bad.docx"
    path.write_bytes(b"not a zip")
    with pytest.raises(LoaderError, match="無法開啟 Word 檔"):
        load(path)


def test_docx_empty(tmp_path):
    path = make_docx(tmp_path / "a.docx", [])
    with pytest.raises(LoaderError, match="沒有可擷取的文字"):
        load(path)


# ---------- Markdown ----------

def test_markdown_heading_path(write_file):
    path = write_file("作品集.md", "\n".join([
        "前言文字",
        "# 作品集",
        "## 小算盤記帳 App",
        "記帳 App。",
        "",
        "",
        "",
        "支援 iOS。",
        "### 技術",
        "Python ##",
        "## 個人部落格 ##",
        "寫技術文章。",
        "#沒有空白不是標題",
    ]))
    assert load(path) == [
        Section(text="前言文字", heading=None),
        Section(text="記帳 App。\n\n支援 iOS。", heading="作品集 > 小算盤記帳 App"),
        Section(text="Python ##", heading="作品集 > 小算盤記帳 App > 技術"),
        Section(text="寫技術文章。\n#沒有空白不是標題", heading="作品集 > 個人部落格"),
    ]


def test_markdown_ignores_headings_in_code_fence(write_file):
    path = write_file("a.md", "# 範例\n```bash\n# 這是註解\n~~~\n```\n說明")
    assert load(path) == [Section(text="```bash\n# 這是註解\n~~~\n```\n說明", heading="範例")]


def test_markdown_big5(tmp_path):
    path = tmp_path / "a.md"
    path.write_bytes("# 標題\n繁體中文內容".encode("cp950"))
    assert load(path) == [Section(text="繁體中文內容", heading="標題")]


def test_markdown_unknown_encoding(tmp_path):
    path = tmp_path / "a.md"
    path.write_bytes(b"\xff\xfe\xff\xfe\x81\x7f")
    with pytest.raises(LoaderError, match="無法辨識文字編碼"):
        load(path)


# ---------- FAQ ----------

FAQ_ROWS = [
    ["question", "answer", "category"],
    ["可以多人共用帳本嗎？", "進階版支援最多 5 人共用。", "功能"],
    ["", "", ""],
    ["", "", "只有分類的列會略過"],
    ["資料會備份嗎？", "每日自動備份。", ""],
]

FAQ_EXPECTED = [
    Section(text="問：可以多人共用帳本嗎？\n答：進階版支援最多 5 人共用。", heading="功能"),
    Section(text="問：資料會備份嗎？\n答：每日自動備份。", heading=None),
]


@pytest.mark.parametrize("encoding", ["utf-8-sig", "utf-8", "cp950"])
def test_faq_csv_encodings(tmp_path, encoding):
    path = make_csv(tmp_path / "faq.csv", FAQ_ROWS, encoding)
    assert load(path) == FAQ_EXPECTED


def test_faq_xlsx(tmp_path):
    path = make_xlsx(tmp_path / "faq.xlsx", FAQ_ROWS)
    assert load(path) == FAQ_EXPECTED


def test_faq_chinese_headers_and_column_order(tmp_path):
    path = make_csv(tmp_path / "faq.csv", [[" 答案 ", "問題"], ["免費。", "要錢嗎？"]])
    assert load(path) == [Section(text="問：要錢嗎？\n答：免費。", heading=None)]


def test_faq_extra_columns_ignored(tmp_path):
    rows = [["分類", "問題", "問題發想原因", "回答"], ["領導", "團隊多大？", "內部備註，不應匯入", "約 40 人。"]]
    path = make_csv(tmp_path / "faq.csv", rows)
    assert load(path) == [Section(text="問：團隊多大？\n答：約 40 人。", heading="領導")]


def test_faq_xlsx_numbers(tmp_path):
    path = tmp_path / "faq.xlsx"
    make_xlsx(path, [["Question", "Answer"], ["進階版多少錢？", 60]])
    assert load(path) == [Section(text="問：進階版多少錢？\n答：60", heading=None)]


def test_faq_missing_column(tmp_path):
    path = make_csv(tmp_path / "faq.csv", [["question", "reply"], ["a", "b"]])
    with pytest.raises(LoaderError, match="缺少欄位 answer"):
        load(path)


def test_faq_incomplete_row(tmp_path):
    path = make_csv(tmp_path / "faq.csv", [["question", "answer"], ["a", "b"], ["只有問題", ""]])
    with pytest.raises(LoaderError, match="第 3 列缺少問題或答案"):
        load(path)


def test_faq_short_row(tmp_path):
    path = tmp_path / "faq.csv"
    path.write_text("question,answer\n只有問題\n", encoding="utf-8")
    with pytest.raises(LoaderError, match="第 2 列"):
        load(path)


def test_faq_header_only(tmp_path):
    path = make_csv(tmp_path / "faq.csv", [["question", "answer"]])
    with pytest.raises(LoaderError, match="沒有可擷取的文字"):
        load(path)


def test_faq_xlsx_broken(tmp_path):
    path = tmp_path / "bad.xlsx"
    path.write_bytes(b"not a zip")
    with pytest.raises(LoaderError, match="無法開啟 Excel 檔"):
        load(path)

